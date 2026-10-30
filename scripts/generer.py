"""Génère la grille Excel complétée et les fichiers de données LaTeX.

Usage : python3 scripts/generer.py
"""
import datetime as dt
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

import donnees as D

RACINE = Path(__file__).resolve().parent.parent
MODELE = RACINE / "scripts" / "modele" / "04_Grille_Audit_SSI.xlsx"
SORTIE_XLSX = RACINE / "livrables" / "MEDICA-SERVICES_Grille_Audit_SSI_completee.xlsx"
DONNEES_TEX = [RACINE / "src" / "rapport" / "donnees", RACINE / "src" / "guide" / "donnees"]

CONSTATS = {c["id"]: c for c in D.CONSTATS}


# ---------------------------------------------------------------------------
# Calculs de contrôle (reproduisent les formules de la grille)
# ---------------------------------------------------------------------------
def calculs():
    maxima, somme, niveau, cible = D.calcul_cible()
    for code, cout in D.COUTS.items():
        assert code in D.ACTIONS, code
    total = sum(D.COUTS[a] * D.ACTIONS[a][2] for a in D.ACTIONS)
    assert total <= D.PLAFOND, "dépassement du plafond"
    return dict(maxima=maxima, somme=somme, niveau=niveau, cible=cible, total=total, marge=D.PLAFOND - total)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def remplir_excel(k):
    wb = openpyxl.load_workbook(MODELE)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws = wb["01_AUTODIAGNOSTIC"]
    for i, (qid, intitule, note, refs, just, conf) in enumerate(D.QUESTIONS, start=2):
        assert ws[f"A{i}"].value == qid and ws[f"B{i}"].value == intitule, (qid, ws[f"B{i}"].value)
        ws[f"C{i}"] = note
        ws[f"D{i}"] = refs
        ws[f"E{i}"] = just
        ws[f"F{i}"] = conf
        for col in "DE":
            ws[f"{col}{i}"].alignment = wrap
    ws["B22"] = D.AJUSTEMENT

    ws = wb["02_MATURITE_EFFECTIVE"]
    for i, p in enumerate(D.PROCESSUS, start=2):
        assert ws[f"A{i}"].value == p["nom"], (p["nom"], ws[f"A{i}"].value)
        ws[f"B{i}"] = p["revendique"]
        ws[f"C{i}"] = p["effectif"]
        ws[f"D{i}"] = p["preuves"]
        ws[f"E{i}"] = p["reserve"]
        ws[f"G{i}"] = p["progression"]
        for col in "DEG":
            ws[f"{col}{i}"].alignment = wrap

    ws = wb["03_CONSTATS_TECHNIQUES"]
    for i, c in enumerate(D.CONSTATS, start=2):
        assert ws[f"A{i}"].value == c["id"]
        valeurs = [c["pieces"], c["observation"], c["reserve"], c["processus"], c["scenario"],
                   c["priorite"], c["reco"], c["cloture"]]
        for col, v in zip("BCDEFGHI", valeurs):
            ws[f"{col}{i}"] = v
            ws[f"{col}{i}"].alignment = wrap

    ws = wb["04_REGISTRE_RISQUES"]
    for i, r in enumerate(D.RISQUES, start=2):
        assert ws[f"A{i}"].value == r["id"]
        valeurs = {"B": r["evenement"], "C": r["actif"], "D": r["menace"], "E": r["vuln"], "F": r["mesures"],
                   "G": r["impact"], "H": r["vrais"], "J": r["decision"], "K": r["preuves"]}
        for col, v in valeurs.items():
            ws[f"{col}{i}"] = v
            if isinstance(v, str):
                ws[f"{col}{i}"].alignment = wrap

    ws = wb["05_PLAN_ACTION"]
    for i in range(2, 13):
        code = ws[f"A{i}"].value
        assert ws[f"D{i}"].value == D.COUTS[code], code
        pilote, echeance, retenu, just, _ = D.ACTIONS[code]
        ws[f"E{i}"] = pilote
        ws[f"F{i}"] = echeance
        ws[f"G{i}"] = retenu
        ws[f"H{i}"] = just
        for col in "FH":
            ws[f"{col}{i}"].alignment = wrap

    # 06_BAREME reste vierge : réservé à l'enseignant.
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    SORTIE_XLSX.parent.mkdir(exist_ok=True)
    wb.save(SORTIE_XLSX)


# ---------------------------------------------------------------------------
# LaTeX
# ---------------------------------------------------------------------------
def tex(s):
    s = str(s)
    for a, b in [("\\", r"\textbackslash{}"), ("&", r"\&"), ("%", r"\%"), ("#", r"\#"), ("_", r"\_"),
                 ("$", r"\$"), ("{", r"\{"), ("}", r"\}")]:
        s = s.replace(a, b)
    return s


def pieces(s):
    return tex(s).replace("/", "/\\allowbreak{}")


def euros(n):
    return f"{n:,}".replace(",", "\\,") + "\\,€"


def defn(nom, valeur):
    return f"\\expandafter\\def\\csname {nom}\\endcsname{{{valeur}}}\n"


def ecrire(nom, contenu):
    for d in DONNEES_TEX:
        d.mkdir(parents=True, exist_ok=True)
        (d / nom).write_text(contenu, encoding="utf-8")


def date(s):
    return dt.date.fromisoformat(s)


def generer_tex(k):
    m = ["% Fichier généré par scripts/generer.py depuis scripts/donnees.py. Ne pas éditer à la main.\n"]
    m.append(f"\\newcommand\\Cible{{{k['cible']}}}\n\\newcommand\\NiveauCalcule{{{k['niveau']}}}\n")
    m.append(f"\\newcommand\\SommeMaxima{{{k['somme']}}}\n\\newcommand\\Ajustement{{{D.AJUSTEMENT}}}\n")
    for (nom, _, _), v, macro in zip(D.GROUPES, k["maxima"], ["MaxConseq", "MaxSensib", "MaxMenaces", "MaxVuln"]):
        m.append(f"\\newcommand\\{macro}{{{v}}}\n")
    m.append(f"\\newcommand\\Plafond{{{euros(D.PLAFOND)}}}\n\\newcommand\\TotalRetenu{{{euros(k['total'])}}}\n")
    m.append(f"\\newcommand\\Marge{{{euros(k['marge'])}}}\n")
    m.append("\\newcommand\\qnote[1]{\\csname qnote#1\\endcsname}\n")
    m.append("\\newcommand\\peff[1]{\\csname peff#1\\endcsname}\n")
    m.append("\\newcommand\\rimp[1]{\\csname rimp#1\\endcsname}\n\\newcommand\\rvra[1]{\\csname rvra#1\\endcsname}\n")
    m.append("\\newcommand\\rscore[1]{\\csname rscore#1\\endcsname}\n\\newcommand\\acout[1]{\\csname acout#1\\endcsname}\n")
    for q in D.QUESTIONS:
        m.append(defn(f"qnote{q[0]}", q[2]))
    for p in D.PROCESSUS:
        m.append(defn(f"peff{p['court']}", p["effectif"]))
    for r in D.RISQUES:
        m.append(defn(f"rimp{r['id']}", r["impact"]) + defn(f"rvra{r['id']}", r["vrais"]))
        m.append(defn(f"rscore{r['id']}", r["impact"] * r["vrais"]))
    for a, c in D.COUTS.items():
        m.append(defn(f"acout{a}", euros(c)))
    ecrire("macros.tex", "".join(m))

    # Tableau Q01 à Q12
    lignes = []
    for qid, intitule, note, refs, just, conf in D.QUESTIONS:
        lignes.append(f"{qid} & {tex(intitule)} & \\textbf{{{note}}} & {tex(refs)} & {tex(just)} & {tex(conf)} \\\\\n")
    ecrire("tab_questions.tex", "".join(lignes))

    # Tableau des maxima
    lignes = []
    for (nom, plage, _), v in zip(D.GROUPES, k["maxima"]):
        lignes.append(f"{tex(nom)} & {plage} & {v} \\\\\n")
    ecrire("tab_maxima.tex", "".join(lignes))

    # Processus
    lignes = []
    for p in D.PROCESSUS:
        ecart = max(0, k["cible"] - p["effectif"])
        lignes.append(f"{tex(p['nom'])} & {tex(p['revendique'])} & \\textbf{{{p['effectif']}}} & {tex(p['suivant'])} & "
                      f"{tex(p['preuves'])} & {ecart} \\\\\n")
    ecrire("tab_processus.tex", "".join(lignes))
    lignes = []
    for p in D.PROCESSUS:
        lignes.append(f"{tex(p['nom'])} & {tex(p['preuves'])} & {tex(p['reserve'])} & {tex(p['progression'])} \\\\\n")
    ecrire("tab_processus_detail.tex", "".join(lignes))
    coords = " ".join(f"({p['effectif']},{i})" for i, p in enumerate(D.PROCESSUS, start=1))
    noms = ",".join(tex(p["court"]) for p in D.PROCESSUS)
    ecrire("graph_maturite.tex", f"\\def\\MaturiteNoms{{{noms}}}\n\\def\\MaturiteNb{{{len(D.PROCESSUS)}}}\n"
           f"\\def\\MaturitePlot{{\\addplot[draw=accent, fill=accent] coordinates {{{coords}}};}}\n")

    # Constats (complet) et corrélation T5
    lignes = []
    for c in D.CONSTATS:
        lignes.append(f"{c['id']} & {pieces(c['pieces'])} & {tex(c['observation'])} & {tex(c['reserve'])} & "
                      f"{tex(c['priorite'])} \\\\\n")
    ecrire("tab_constats.tex", "".join(lignes))
    lignes = []
    for c in D.CONSTATS:
        lignes.append(f"{c['id']} & {tex(c['processus'])} & {tex(c['scenario'])} & {tex(c['reco'])} & {tex(c['cloture'])} \\\\\n")
    ecrire("tab_constats_suite.tex", "".join(lignes))
    lignes = []
    for cid in D.CORRELATIONS:
        c = CONSTATS[cid]
        lignes.append(f"{cid} {tex(c['titre'])} & {pieces(c['pieces'])} & {c['risque']} & {tex(c['processus'])} & "
                      f"{tex(c['niveau'])} & {tex(c['action'])} & {tex(c['pilote'])} & {tex(c['cloture'])} \\\\\n")
    ecrire("tab_correlation.tex", "".join(lignes))

    # Risques
    lignes = []
    for r in D.RISQUES:
        lignes.append(f"{r['id']} & {tex(r['evenement'])} & {tex(r['actif'])} & {tex(r['vuln'])} & {r['impact']} & {r['vrais']} & "
                      f"\\textbf{{{r['impact'] * r['vrais']}}} & {tex(r['decision'])} \\\\\n")
    ecrire("tab_risques.tex", "".join(lignes))
    lignes = []
    for r in D.RISQUES:
        lignes.append(f"{r['id']} & {tex(r['menace'])} & {tex(r['mesures'])} & {tex(r['preuves'])} \\\\\n")
    ecrire("tab_risques_suite.tex", "".join(lignes))
    cases = {}
    for r in D.RISQUES:
        cases.setdefault((r["impact"], r["vrais"]), []).append(r["id"])
    pts = "".join(f"\\RisqueCase{{{v}}}{{{i}}}{{{', '.join(ids)}}}\n" for (i, v), ids in cases.items())
    ecrire("graph_risques.tex", pts)

    # Actions
    lignes = []
    for code in D.COUTS:
        pilote, echeance, retenu, just, court = D.ACTIONS[code]
        etat = "\\Retenu" if retenu else "\\Differe"
        lignes.append(f"{code} & {tex(court)} & {euros(D.COUTS[code])} & {etat} & {tex(pilote)} & {tex(echeance)} & {tex(just)} \\\\\n")
    ecrire("tab_actions.tex", "".join(lignes))

    # Feuille de route SSI (12 mois)
    lignes = []
    for libelle, debut, fin, livrable in D.FEUILLE_ROUTE:
        lignes.append(f"\\FeuilleRoute{{{tex(libelle)}}}{{{debut + 1}}}{{{fin}}}{{{tex(livrable)}}}\n")
    ecrire("gantt_feuille_route.tex", "".join(lignes))
    lignes = []
    for nom, (m6, m12) in D.TRAJECTOIRE.items():
        eff = next(p["effectif"] for p in D.PROCESSUS if p["court"] == nom)
        lignes.append(f"{tex(nom)} & {eff} & {m6} & {m12} & {k['cible']} \\\\\n")
    ecrire("tab_trajectoire.tex", "".join(lignes))

    # Couleurs des responsables
    ecrire("couleurs.tex", "".join(f"\\definecolor{{resp{n}}}{{HTML}}{{{c}}}\n" for n, c in D.COULEURS.items()))

    # Gantt du projet (groupé par phase, couleur par responsable, liens fin-début)
    debut = date("2026-10-20")
    jour = lambda d: (date(d) - debut).days + 1
    lignes = []
    for phase in D.PHASES:
        taches = [t for t in D.TACHES if t[6] == phase]
        lignes.append(f"\\PhaseGantt{{{tex(phase)}}}{{{min(jour(t[3]) for t in taches)}}}{{{max(jour(t[4]) for t in taches)}}}\n")
        for tid, lib, resp, d0, d1, deps, _ in taches:
            lignes.append(f"\\TacheGantt{{{tid}}}{{{tex(lib)}}}{{{tex(resp)}}}{{{jour(d0)}}}{{{jour(d1)}}}\n")
    for jid, lib, d in D.JALONS:
        lignes.append(f"\\JalonGantt{{{jid}}}{{{tex(lib)}}}{{{jour(d)}}}\n")
    liens = "".join(f"\\ganttlink{{{a_}}}{{{b_}}}\n" for a_, b_ in D.LIENS_GANTT)
    ecrire("gantt_projet.tex", "".join(lignes) + liens + f"\\def\\JourKanban{{{jour(D.DATE_KANBAN)}}}\n")

    # Kanban : état attesté
    colonnes = ["À faire", "En cours", "À relire", "À corriger", "Validé"]
    f = lambda s: date(s).strftime("%d/%m")
    lignes = []
    for tid, lib, resp, d0, d1, deps, _ in D.TACHES:
        col, note = D.KANBAN[tid]
        lignes.append(f"\\CarteKanban{{{colonnes.index(col)}}}{{{tid}}}{{{tex(lib)}}}{{{tex(resp)}}}{{{f(d0)}–{f(d1)}}}{{{tex(note)}}}\n")
    compte = {c: sum(1 for v in D.KANBAN.values() if v[0] == c) for c in colonnes}
    ecrire("kanban.tex", "".join(lignes))
    ecrire("kanban_comptes.tex", "\\def\\KanbanComptes{" + ",".join(f"{tex(c)}/{n}" for c, n in compte.items()) + "}\n")

    # Tableau des tâches (annexe)
    lignes = []
    for tid, lib, resp, d0, d1, deps, phase in D.TACHES:
        col, note = D.KANBAN[tid]
        etat = tex(col) + (f" ({tex(note)})" if note else "")
        lignes.append(f"{tid} & {tex(lib)} & {tex(resp)} & {f(d0)} au {f(d1)} & {', '.join(deps) or '·'} & {etat} \\\\\n")
    ecrire("tab_taches.tex", "".join(lignes))

    # Tableaux courts du corps du rapport
    ecrire("tab_questions_court.tex", "".join(
        f"{q[0]} & {tex(q[1])} & \\textbf{{{q[2]}}} & {tex(q[3])} & {tex(q[5])} \\\\\n" for q in D.QUESTIONS))
    ecrire("tab_risques_court.tex", "".join(
        f"{r['id']} & {tex(r['court'])} & {tex(r['actif'])} & {r['impact']} & {r['vrais']} & \\textbf{{{r['impact'] * r['vrais']}}} & {tex(D.DECISION_COURTE[r['id']])} \\\\\n"
        for r in D.RISQUES))
    ecrire("tab_risques_court_sans_actif.tex", "".join(
        f"{r['id']} & {tex(r['court'])} & {r['impact']} & {r['vrais']} & \\textbf{{{r['impact'] * r['vrais']}}} & {tex(D.DECISION_COURTE[r['id']])} \\\\\n"
        for r in D.RISQUES))
    ecrire("tab_actions_court.tex", "".join(
        f"{c} & {tex(D.ACTIONS[c][4])} & {euros(D.COUTS[c])} & {'\\Retenu' if D.ACTIONS[c][2] else '\\Differe'} & {tex(D.ACTIONS[c][0])} \\\\\n"
        for c in D.COUTS))

    # Données des graphiques
    g = []
    g.append("\\def\\NotesQ{" + ",".join(f"{q[0]}/{q[2]}" for q in D.QUESTIONS) + "}\n")
    # grille des exigences et cohérence avec les niveaux
    for i, p in enumerate(D.PROCESSUS):
        niv = 0
        for n in (1, 2, 3):
            ok = all(e[2][i] in "OD" for e in D.EXIGENCES if e[0] == n)
            if ok and niv == n - 1:
                niv = n
        assert niv == p["effectif"], (p["court"], niv, p["effectif"])
    g.append("\\def\\ProcCourts{" + ",".join(D.PROC_COURTS) + "}\n")
    g.append("\\def\\Exigences{" + ",".join(f"{e[0]}/{tex(e[1])}/{e[2]}" for e in D.EXIGENCES) + "}\n")
    g.append("\\def\\Effectifs{" + ",".join(str(p["effectif"]) for p in D.PROCESSUS) + "}\n")
    g.append("\\def\\TrajM{" + ",".join(f"{D.TRAJECTOIRE[n][0]}/{D.TRAJECTOIRE[n][1]}" for n in D.PROC_COURTS) + "}\n")
    g.append("\\def\\Comptes{" + ",".join(f"{c}/{m}" for c, m in D.COMPTES) + "}\n")
    g.append("\\def\\Sauvegardes{" + ",".join(f"{{{n}}}/{{{d if d else '-'}}}" for n, d in D.SAUVEGARDES) + "}\n")
    g.append("\\def\\Vulns{" + ",".join(f"{a_}/{c}/{r}/{e}/{pr}" for a_, c, r, e, pr in D.VULNS) + "}\n")
    retenues = [c for c in D.COUTS if D.ACTIONS[c][2]]
    differees = [c for c in D.COUTS if not D.ACTIONS[c][2]]
    g.append("\\def\\BudgetRetenu{" + ",".join(f"{c}/{D.COUTS[c] // 1000}/{tex(D.ACTIONS[c][4])}" for c in retenues) + "}\n")
    g.append("\\def\\BudgetDiffere{" + ",".join(f"{c}/{D.COUTS[c] // 1000}/{tex(D.ACTIONS[c][4])}" for c in differees) + "}\n")
    g.append(f"\\def\\BudgetK{{{k['total'] // 1000}}}\\def\\MargeK{{{k['marge'] // 1000}}}"
             f"\\def\\DiffereK{{{sum(D.COUTS[c] for c in differees) // 1000}}}\\def\\PlafondK{{{D.PLAFOND // 1000}}}\n")
    corr = []
    for cid in D.CORRELATIONS:
        c = CONSTATS[cid]
        corr.append(f"\\ChaineCorr{{{cid}}}{{{tex(c['titre'])}}}{{{c['risque']}}}{{{tex(c['processus'])}}}{{{tex(c['niveau'])}}}{{{tex(c['action'])}}}{{{tex(c['pilote'])}}}\n")
    ecrire("graph_correlation.tex", "".join(corr))
    ecrire("graphes.tex", "".join(g))


def main():
    k = calculs()
    if not MODELE.exists():
        raise SystemExit(f"Modèle introuvable : {MODELE}")
    remplir_excel(k)
    generer_tex(k)
    print(f"Cible {k['cible']} (somme {k['somme']}, maxima {k['maxima']}), engagement {k['total']} €, marge {k['marge']} €")


if __name__ == "__main__":
    main()
